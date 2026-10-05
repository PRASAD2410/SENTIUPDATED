import unittest
from io import BytesIO
from datetime import datetime
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from app.file_reader import parse_upload, read_upload, UnsupportedFormatError

class ReaderTests(unittest.TestCase):
    def test_csv_identifiers_and_multiline(self):
        result = parse_upload('calls.csv', b'phone,phone,note\r\n00123,00456,"first\nsecond"\r\n00999,00888,end\r\n')
        self.assertEqual(result['units'][1]['cells'], ['00123', '00456', 'first\nsecond'])
        self.assertEqual(result['units'][2]['line_number'], 4)
        self.assertEqual(result['units'][0]['cells'], ['phone', 'phone', 'note'])

    def test_semicolon_and_tsv(self):
        for name, content, separator in [('a.csv', b'id;amount\n001;0\n', ';'), ('a.tsv', b'id\tamount\n001\t0\n', '\t')]:
            result = parse_upload(name, content)
            self.assertEqual(result['metadata']['delimiter'], separator)
            self.assertEqual(result['units'][1]['cells'], ['001', '0'])

    def test_encodings_and_compatibility(self):
        self.assertEqual(read_upload('a.txt', 'Hindi: नमस्ते'.encode('utf-16')), 'Hindi: नमस्ते')
        result = parse_upload('a.txt', 'café'.encode('cp1252'))
        self.assertEqual(result['text'], 'café')
        self.assertTrue(result['warnings'])

    def test_docx_table_order(self):
        document = Document()
        document.add_paragraph('Before')
        document.add_table(rows=1, cols=2).rows[0].cells[0].text = '00123'
        document.add_paragraph('After')
        output = BytesIO()
        document.save(output)
        result = parse_upload('a.docx', output.getvalue())
        self.assertEqual([u['kind'] for u in result['units']], ['paragraph', 'row', 'paragraph'])
        self.assertEqual(result['text'], 'Before\n00123 | \nAfter')

    def test_xlsx_zero_dates_formulas_and_sheets(self):
        workbook = Workbook()
        workbook.active.title = 'Transfers'
        workbook.active.append(['00123', 0, datetime(2026, 10, 5, 10, 30), '=1+1'])
        workbook.create_sheet('Calls').append(['00456', False])
        output = BytesIO()
        workbook.save(output)
        result = parse_upload('a.xlsx', output.getvalue())
        self.assertEqual(result['units'][0]['cells'], ['00123', '0', '2026-10-05T10:30:00', '=1+1'])
        self.assertEqual(result['units'][1]['sheet'], 'Calls')
        self.assertIn('False', result['text'])
        self.assertTrue(any('Formula' in warning for warning in result['warnings']))

    def test_blank_and_locked_pdf(self):
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        output = BytesIO()
        writer.write(output)
        result = parse_upload('a.pdf', output.getvalue())
        self.assertEqual(result['text'], '')
        self.assertEqual(result['units'][0]['page_number'], 1)
        self.assertTrue(any('OCR' in warning for warning in result['warnings']))
        writer.encrypt('secret')
        output = BytesIO()
        writer.write(output)
        with self.assertRaisesRegex(ValueError, 'Password-protected'):
            parse_upload('a.pdf', output.getvalue())

    def test_empty_and_unsupported(self):
        self.assertEqual(parse_upload('a.csv', b',,\n,,\n')['text'], '')
        with self.assertRaises(UnsupportedFormatError):
            parse_upload('a.exe', b'123')

if __name__ == '__main__':
    unittest.main()
